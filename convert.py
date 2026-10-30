"""Konversi data angkatan (Excel / CSV / Google Sheets) ke data/mahasiswa.json.

Pemakaian:
    python convert.py data_angkatan.xlsx      # file Excel (butuh: pip install openpyxl)
    python convert.py data_angkatan.csv       # file CSV
    python convert.py                         # unduh langsung dari Google Sheets
                                              # (sheet harus "Anyone with the link can view")

Baris header dicari otomatis: baris pertama yang memuat kolom NAMA dan NIM.
"""
import csv
import io
import json
import sys
import urllib.request
from pathlib import Path

SHEET_ID = "10p-XeUAmxi-3J2ljUYxUkGuCps8SN5GT6paaVAZ_T4A"
OUTPUT = Path(__file__).parent / "data" / "mahasiswa.json"
FOTO_DIR = Path(__file__).parent / "img" / "foto"


def read_rows(source):
    if source is None:
        url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=csv&gid=0"
        with urllib.request.urlopen(url) as r:
            return list(csv.reader(io.StringIO(r.read().decode("utf-8"))))
    path = Path(source)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        return [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.reader(f))


def to_records(rows):
    for i, row in enumerate(rows):
        header = [c.strip().upper() for c in row]
        if "NAMA" in header and "NIM" in header:
            col_nama, col_nim = header.index("NAMA"), header.index("NIM")
            col_no = header.index("NO") if "NO" in header else None
            break
    else:
        sys.exit("Header NAMA/NIM tidak ditemukan.")

    records = []
    for row in rows[i + 1:]:
        cell = lambda c: row[c].strip() if c is not None and c < len(row) else ""
        nama, nim = " ".join(cell(col_nama).split()), cell(col_nim).upper()
        if not nama or not nim:
            continue
        no = cell(col_no)
        record = {
            "no": int(float(no)) if no else len(records) + 1,
            "nama": nama,
            "nim": nim,
        }
        # Foto opsional: img/foto/<NIM>.jpg
        if (FOTO_DIR / f"{nim}.jpg").exists():
            record["foto"] = f"img/foto/{nim}.jpg"
        records.append(record)
    return records


if __name__ == "__main__":
    records = to_records(read_rows(sys.argv[1] if len(sys.argv) > 1 else None))
    OUTPUT.parent.mkdir(exist_ok=True)
    OUTPUT.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{len(records)} mahasiswa ditulis ke {OUTPUT}")
